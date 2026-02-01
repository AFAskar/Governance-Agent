"""
DOCX parser: extract text from Word documents.
"""

from pathlib import Path


def extract_text_from_docx(path: str) -> str:
    """
    Extract text from a DOCX file. Reads paragraphs and table cells.
    """
    from docx import Document

    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")
    try:
        doc = Document(path)
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n\n".join(parts)
    except Exception as e:
        raise ValueError(f"Could not read DOCX {path}: {e}") from e
