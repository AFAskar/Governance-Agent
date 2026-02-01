"""
PPTX parser: extract text from PowerPoint files.
"""

from pathlib import Path


def extract_text_from_pptx(path: str) -> str:
    """
    Extract text from a PPTX file. Iterates slides and shape text.
    """
    from pptx import Presentation

    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"PPTX file not found: {path}")
    try:
        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text.strip())
        return "\n\n".join(p for p in parts if p)
    except Exception as e:
        raise ValueError(f"Could not read PPTX {path}: {e}") from e
